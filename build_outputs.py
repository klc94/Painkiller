#!/usr/bin/env python3
"""Build pain_points.csv, scores.csv and painkiller_workbook.xlsx from reviewed evidence.

Every quote is extracted verbatim from the pulled post/comment text and asserted to be a
substring of it (whitespace/quote-normalised). Evidence rows are ones I read and judged relevant;
keyword matches I did not review are NOT included.
"""
import csv, glob, json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAW = "reddit_raw"
TR = str.maketrans({"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u00a0": " "})


def norm(s):
    return " ".join(s.translate(TR).split())


POSTS = {}
for path in glob.glob(f"{RAW}/*/posts_year.jsonl"):
    for line in open(path):
        p = json.loads(line)
        POSTS[p["id"]] = p


def walk(n, out):
    if isinstance(n, dict):
        if "body" in n and "id" in n and n.get("body") not in ("[deleted]", "[removed]"):
            out[n["id"]] = n
        for v in n.values():
            if isinstance(v, (dict, list)):
                walk(v, out)
    elif isinstance(n, list):
        for x in n:
            walk(x, out)


COMMENTS = {}
for f in glob.glob(f"{RAW}/*/comments/*.json"):
    try:
        d = json.load(open(f))
    except Exception:
        continue
    tmp = {}
    walk(d, tmp)
    for cid, c in tmp.items():
        c["_post"] = os.path.basename(f)[:-5]
        c["_sub"] = f.split("/")[1]
        COMMENTS[cid] = c

MONEY = re.compile(r"\$\s?\d[\d,]*(?:\.\d+)?\s?[kKmM]?(?:\s?(?:/|per|a)\s?(?:mo|month|yr|year|hour|hr))?|\b\d+\s?%|\b\d+\s+(?:hours?|hrs?|minutes?|days?|weeks?|months?)\b")


def window(text, phrase, maxw=24):
    t, p = norm(text), norm(phrase)
    i = t.lower().find(p.lower())
    if i < 0:
        return None
    j = i + len(p)
    while i > 0 and t[i - 1] != " ":
        i -= 1
    while j < len(t) and t[j] != " ":
        j += 1
    span = t[i:j].split()
    if len(span) > maxw:
        span = span[:maxw]
    room = maxw - len(span)
    before = t[:i].split()
    before = before[-(room // 2):] if room // 2 else []
    after = t[j:].split()[:room - len(before)]
    q = " ".join(before + span + after)
    assert q in t, (q, t[:80])
    return q


# (cluster, kind, id, phrase, industry)   kind: p=post, c=comment
E = [
 # --- AR / late payment
 ("B Overdue invoices", "p", "1uw84p0", "customer owes me $14k and has stopped replying", "services/contracting"),
 ("B Overdue invoices", "p", "1ovpacl", "ghosted me and have stopped trying to pay anything at all", "services"),
 ("B Overdue invoices", "p", "1o8efjo", "Worth it to sue a customer for $14,000 in unpaid invoices?", "services"),
 ("B Overdue invoices", "p", "1sh4b89", "chasing payments", "services"),
 ("B Overdue invoices", "p", "1t4hryr", "How much time do you guys waste chasing unpaid invoices?", "services"),
 ("B Overdue invoices", "p", "1p0zi4o", "Wholesale account owes me $25K and keeps ghosting me", "wholesale/apparel"),
 ("B Overdue invoices", "p", "1qzkjmj", "Commercial client not paying", "construction"),
 ("B Overdue invoices", "p", "1sl6m0a", "somewhere around 180k sitting in unpaid invoices", "B2B services"),
 ("B Overdue invoices", "p", "1pufmg5", "Racked up $70K of legal fees", "contracting"),
 ("B Overdue invoices", "p", "1q068o1", "How many follow-ups before you give up", "contracting"),
 ("B Overdue invoices", "p", "1s87s5s", "I'm good at pulling wire but I'm terrible at chasing invoices at 9pm", "electrical"),
 ("B Overdue invoices", "p", "1s51cf8", "sales rep told customer they could have net 75 now customer won't pay at net 30", "B2B (accounting view)"),
 ("B Overdue invoices", "p", "1ppx9pp", "Finally got paid $15k from a GC who ghosted me for 4 months", "electrical contracting"),
 ("B Overdue invoices", "p", "1qtc2o6", "I still have thousands of dollars of unpaid invoices", "law"),
 ("B Overdue invoices", "p", "1pp11hg", "My unpaid balances are up slightly to $35,000", "law"),
 ("B Overdue invoices", "p", "1uyg6kr", "We're getting very tired of people not paying their retainers", "law"),
 ("B Overdue invoices", "p", "1owd6da", "Probably going to just write off this $16k bill", "law"),
 ("B Overdue invoices", "c", "oxh6zvz", "I use a collection agency", "services (WTP signal: pays % to collector)"),
 ("B Overdue invoices", "c", "njua7p7", "There's a thing called \"factoring\"", "services (alternative to service)"),
 # --- chargebacks / holds
 ("C Chargebacks and payment holds", "p", "1phlsw0", "Stripe is holding $8,000 of my money indefinitely", "small merchant"),
 ("C Chargebacks and payment holds", "p", "1ruedgx", "Stripe holding $1.2M of our operating funds for 6+ weeks", "merchant"),
 ("C Chargebacks and payment holds", "p", "1qci3a4", "Woke up this morning to seven chargebacks", "ecommerce"),
 ("C Chargebacks and payment holds", "p", "1qci3a4", "Just got charged back $3,400 in one day", "ecommerce"),
 ("C Chargebacks and payment holds", "p", "1rjivlb", "Chargebacks are basically legalized theft", "ecommerce"),
 ("C Chargebacks and payment holds", "p", "1tnt780", "I'm so tired of chargebacks", "events (15,000 attendees)"),
 ("C Chargebacks and payment holds", "p", "1qlt8xl", "they locked my store for 48 hrs becz of a chargeback on a $47 order", "ecommerce"),
 ("C Chargebacks and payment holds", "p", "1rppa5u", "Chargebacks about to sink my whole side hustle", "ecommerce"),
 ("C Chargebacks and payment holds", "p", "1smq52b", "Being scammed with online orders", "restaurant (delivery apps)"),
 ("C Chargebacks and payment holds", "c", "onwo02o", "They say they fight the chargeback on our behalf but who knows if they do", "events (WTP signal: relies on paid tool, distrusts it)"),
 # --- processing fees
 ("G Card processing fee audit (existing)", "p", "1ol8edm", "Just ran the numbers on what payment processing fees actually cost us last year", "small business"),
 ("G Card processing fee audit (existing)", "p", "1qfuedn", "My processing fees have always been around 4%. In 2025 the fees were around 9%", "small business"),
 ("G Card processing fee audit (existing)", "p", "1rathu6", "My biggest surprise expense, which I was aware of but just didnt fully grasp, is credit card fees", "law (solo)"),
 ("G Card processing fee audit (existing)", "c", "nmgclee", "payment processing is the smallest", "counter-evidence: commenter says fees are ~2.5% of gross"),
 ("G Card processing fee audit (existing)", "c", "nmgdsgg", "I owned a restaurant for 10 years. It was my second highest expense.", "restaurant"),
 # --- lost billable time
 ("D Lost billable time (law/consulting)", "p", "1qsqs25", "lose the 0.1 billable increment", "law"),
 ("D Lost billable time (law/consulting)", "p", "1s9nc77", "ended up leaving a bunch of days unbilled", "law"),
 ("D Lost billable time (law/consulting)", "p", "1tan9zr", "Trevor has double billed the same item on the same day", "law"),
 ("D Lost billable time (law/consulting)", "p", "1pkt9t4", "I can't prove how long things took", "electrical"),
 ("D Lost billable time (law/consulting)", "c", "o2xj1r3", "I use Legalmate, it automatically logs all of my phone calls and texts", "law (counter-evidence: AI tools already exist)"),
 # --- intake / no-shows
 ("E Missed calls, intake, no-shows", "p", "1pmhgi9", "I just spent my Saturday afternoon handling intake for three potential clients", "law (solo)"),
 ("E Missed calls, intake, no-shows", "p", "1rm6jq3", "we're talking lost revenue, wasted chair time, staff standing around", "dental"),
 ("E Missed calls, intake, no-shows", "p", "1r1nfib", "My front desk person is juggling phones, patient intake, appointment scheduling, referral coordination, and prior authorizations", "dermatology"),
 # --- denials / clinics
 ("A Small-dollar payer denials and underpayments", "p", "1tbdqtt", "it takes a biller (paid \\~$28/hr) about 45 minutes to an hour to research the denial", "physical therapy"),
 ("A Small-dollar payer denials and underpayments", "p", "1tbdqtt", "They deny small-balance claims for nonsense reasons, knowing your staff will never chase them.", "physical therapy (author also sells a tool: self-interested)"),
 ("A Small-dollar payer denials and underpayments", "p", "1r6fht0", "Insurance denied my patient's 7th visit even though she still can't climb stairs", "physical therapy"),
 ("A Small-dollar payer denials and underpayments", "p", "1turokv", "A ridiculous workers comp denial", "physical therapy"),
 ("A Small-dollar payer denials and underpayments", "p", "1w7etcs", "I pay a billing company $1,000 month", "chiropractic (WTP signal)"),
 ("A Small-dollar payer denials and underpayments", "c", "oir0wft", "I pay for a billing company to do everything", "chiropractic (WTP signal)"),
 ("A Small-dollar payer denials and underpayments", "c", "p7w5t5y", "they only do $300 base pay with 4% collections", "chiropractic (competitor price point)"),
 ("A Small-dollar payer denials and underpayments", "c", "olgdgva", "How HIPAA compliant is it?", "physical therapy (objection: HIPAA)"),
 ("A Small-dollar payer denials and underpayments", "p", "1pvkex8", "get lowballed", "dental (fee schedules: structural, not recoverable)"),
 ("A Small-dollar payer denials and underpayments", "p", "1wo9oi7", "after insurance reimbursement the total net collections was just at about break even", "dental"),
 # --- COI / compliance
 ("F Vendor COI / prequal compliance", "p", "1sh61hu", "I swear I spend half my Mondays just staring at a spreadsheet trying to figure out whose COI is about to expire", "landscaping/property services"),
 ("F Vendor COI / prequal compliance", "p", "1tdbch2", "lost a 100k contract because my safety paperwork wasn't up to their standard", "landscaping/civil"),
 ("F Vendor COI / prequal compliance", "p", "1wellfk", "Custom COI requests with multiple additional insureds", "services to enterprise"),
 # --- M trucking detention claims
 ("M Detention / accessorial claims for small carriers", "p", "1rjuuq6", "287 hours of detention last year. At $75/hr after free time that's $21,525 just gone", "trucking (8-truck fleet; unverified, low-engagement post)"),
 ("M Detention / accessorial claims for small carriers", "p", "1rdilf4", "The carrier eats 4 hours of detention because filing is a hassle and the broker knows it", "trucking (owner-operator)"),
 ("M Detention / accessorial claims for small carriers", "p", "1p3faoc", "If ELDs track your truck, why do brokers still deny detention pay?", "trucking (post may be idea-validation by a builder)"),
 ("M Detention / accessorial claims for small carriers", "p", "1nz38sy", "won't pay detention pay until after 2 hours", "trucking (company driver)"),
 ("M Detention / accessorial claims for small carriers", "p", "1w3uhts", "Checked in at 9:30am, now 6pm warehouse closed for the day", "trucking (driver)"),
 ("M Detention / accessorial claims for small carriers", "p", "1qofr8a", "refuses to acknowledge the $1000 in missing payments, including detentions", "trucking (driver vs carrier)"),
 ("M Detention / accessorial claims for small carriers", "p", "1v4vfhz", "What can I do if broker is refusing to pay layover?", "trucking"),
 ("M Detention / accessorial claims for small carriers", "p", "1qj8c52", "got a 500 deduction for missing an appt time by 1 minute", "trucking"),
 ("M Detention / accessorial claims for small carriers", "p", "1p8mk2x", "then his company refused to pay the full amount", "trucking (owner-operator vs broker)"),
 ("M Detention / accessorial claims for small carriers", "p", "1q25nof", "$17/hr Detention that only starts 2hrs after the appointment time", "trucking (driver)"),
 ("M Detention / accessorial claims for small carriers", "p", "1qqe635", "Carrier is demanding a layover from Tuesday to Wednesday. Is compensation owed to the carrier in this situation?", "freight brokerage (broker side: claims are judgment calls)"),
 # --- N unbilled extras
 ("N Unbilled extras / change orders (contractors)", "p", "1qnua1c", "I gave away about 4 hours a week. At my rate, that is $12,000 a year", "contracting"),
 ("N Unbilled extras / change orders (contractors)", "p", "1u8u7wr", "any additional work would be handled through change orders and paid separately", "contracting"),
 ("N Unbilled extras / change orders (contractors)", "p", "1nvziy3", "My biggest client asked for a week of free work", "services"),
 ("M Detention / accessorial claims for small carriers", "c", "o8g9h0p", "I think most people would pay a fee if you were able to recover this money for them", "trucking (WTP signal, one commenter)"),
 ("M Detention / accessorial claims for small carriers", "c", "o51pqv2", "The minute you file a claim on a brokers bond, they DNU", "trucking (counter-evidence: retaliation risk)"),
 ("M Detention / accessorial claims for small carriers", "c", "o76ubzo", "Every broker handles it differently, every shipper has different rules", "trucking (counter-evidence: hard to automate)"),
 ("M Detention / accessorial claims for small carriers", "c", "nq4bkbi", "Brokers pocket it.", "trucking"),
 ("N Unbilled extras / change orders (contractors)", "c", "o1whurk", "Just raise your estimate by a percentage to cover the", "contracting (counter-evidence: the fix is pricing, not a service)"),
 # --- existing candidates (thin)
 ("J HVAC warranty recovery (existing)", "p", "1o2h1lw", "Impossible to make a profit recently with all the free warranty labor work from shit products", "HVAC"),
 ("K Delivery-app disputes (existing)", "p", "1o72aqi", "quietly stacks and double-charges", "restaurant"),
 ("K Delivery-app disputes (existing)", "p", "1sukdal", "The prices on my page are lower than UberEats because I don't have to inflate them to cover the fees", "restaurant"),
 ("L UPS/FedEx refunds (existing)", "p", "1p68wjj", "UPS is charging us heavy surcharges every week.", "small shipper"),
]


def build_rows():
    rows, bad = [], []
    for cluster, kind, i, phrase, ind in E:
        if kind == "p":
            p = POSTS.get(i)
            if not p:
                bad.append((i, "post missing")); continue
            text, sub = f"{p['title']} {p.get('selftext') or ''}", p["subreddit"]
            url, sc, nc = p["permalink"], p.get("score"), p.get("num_comments")
        else:
            c = COMMENTS.get(i)
            if not c:
                bad.append((i, "comment missing")); continue
            text, sub = c["body"], c["_sub"]
            url = f"https://www.reddit.com/r/{sub}/comments/{c['_post']}/_/{i}/"
            sc, nc = c.get("score"), ""
        q = window(text, phrase)
        if not q:
            bad.append((i, "phrase not found: " + phrase[:50])); continue
        imp = "; ".join(dict.fromkeys(m.group(0).strip() for m in MONEY.finditer(q)))
        if not imp:  # fall back to post text
            imp = "; ".join(list(dict.fromkeys(m.group(0).strip() for m in MONEY.finditer(norm(text)[:600])))[:3])
        rows.append([cluster, sub, url, q, ind, imp, sc, nc, "post" if kind == "p" else "comment"])
    return rows, bad


# ---------------------------------------------------------------- scoring
CRIT = ["1 Painkiller strength", "2 Frequency", "3 Willingness to pay", "4 AI resistance (3yr)",
        "5 Delivery automatable", "6 Outreach automatable (x2)", "7 Massive pain", "8 Purchasing power",
        "9 Easy to target", "10 Growing", "11 Dream outcome", "12 Perceived likelihood (provable upfront)",
        "13 Speed to result", "14 Low client effort", "15 Grand Slam Offer potential", "16 Recurring revenue",
        "17 LTV:CAC", "18 Gross margin at scale", "19 Reachable by paid ads + direct mail"]
CL = {
 "A Small-dollar payer denials and underpayments (independent clinics)": ([4,3,4,4,3,4,3,3,5,4,4,5,3,3,5,4,4,3,5], "NEW"),
 "B Overdue-invoice follow-up + lien/notice deadlines (owner-name, not collections)": ([4,5,2,2,4,4,4,3,4,3,3,3,4,3,3,4,2,4,4], "NEW"),
 "C Chargeback / processor-hold dispute packaging": ([4,4,3,3,4,4,4,2,3,4,3,4,3,4,5,4,3,4,4], "NEW"),
 "D Lost billable time capture (small law/consulting)": ([3,3,3,1,4,5,2,4,5,3,3,4,4,3,3,5,3,5,5], "NEW"),
 "E Missed calls / intake / no-show follow-up": ([4,5,4,2,5,4,3,3,4,3,3,3,5,4,3,5,2,4,5], "NEW (crowded)"),
 "F Vendor COI / prequal compliance tracking": ([3,2,2,3,4,4,2,3,4,3,3,2,3,3,2,4,2,4,4], "NEW"),
 "M Detention / accessorial claim filing for small carriers": ([4,3,3,4,3,5,3,2,5,3,3,4,3,3,5,4,3,4,5], "NEW (added after first pass)"),
 "N Unbilled extras / change-order capture (contractors)": ([3,3,2,2,4,4,2,3,4,3,3,3,4,3,2,4,2,5,4], "NEW (added after first pass)"),
 "G Card processing fee audit": ([3,3,2,4,4,4,2,3,4,3,3,5,3,4,4,3,3,4,4], "EXISTING #1"),
 "H Utility bill audit (audit only)": ([3,1,2,4,3,4,1,3,3,2,3,5,2,3,4,2,2,4,3], "EXISTING #2"),
 "I Manufacturer co-op fund recovery": ([3,1,1,4,3,3,1,3,2,2,3,3,2,2,4,3,3,4,3], "EXISTING #3"),
 "J HVAC manufacturer warranty claim recovery": ([3,2,2,4,3,4,2,3,4,2,3,3,3,2,4,3,3,4,4], "EXISTING #4"),
 "K Delivery-app dispute recovery": ([2,2,2,3,4,4,2,2,3,3,2,4,3,3,4,3,2,4,4], "EXISTING #5"),
 "L UPS/FedEx late-delivery refunds": ([2,2,3,4,5,4,2,3,3,2,3,5,4,4,5,3,3,5,3], "EXISTING #6"),
}


def total(v):
    return sum(v) + v[5]  # outreach counted twice; max 100


# $1M math: name -> (avg revenue/client/yr, source of that number, market size, market source)
MATH = {
 "A": (3600, "ASSUMPTION: clinic collects $1.2M x 1.5% recovered x 20% fee", 113000, "DATA: ~33,000 independent PT clinics + ~80,000 independent dental practices (web search); chiro/medical excluded"),
 "B": (1788, "ASSUMPTION: $149/mo done-for-you follow-up", 1050000, "DATA: ~598k specialty-trade establishments + ~450k law firms (web search)"),
 "C": (1500, "ASSUMPTION: $99/mo + share of recovered", 150000, "DATA: 2.5-3.5M live Shopify stores; ASSUMPTION 5% big enough to matter"),
 "D": (1188, "ASSUMPTION: $99/mo", 180000, "DATA: ~450k law firms, ~40% solo"),
 "E": (2388, "ASSUMPTION: $199/mo", 1050000, "DATA: trades + law (web search)"),
 "F": (1188, "ASSUMPTION: $99/mo", 598000, "DATA: ~598k specialty-trade establishments"),
 "M": (2400, "ASSUMPTION: 5-truck carrier, ~$8k/yr unclaimed detention x 30% fee", 500000, "UNVERIFIED ASSUMPTION: ~500k small for-hire carriers. WEB: 2.12M total registered carriers (incl. private fleets); ~91.5% run 10 trucks or fewer"),
 "N": (1188, "ASSUMPTION: $99/mo", 598000, "DATA: ~598k specialty-trade establishments"),
 "G": (750, "ASSUMPTION: $500k card volume x 0.5% saved x ~30% fee + residual", 4500000, "DATA: 6.4M employer firms (SBA); ASSUMPTION ~70% take cards"),
 "H": (450, "ASSUMPTION: $30k utility spend x 5% recovered x 30% fee", 2000000, "ASSUMPTION: multi-site/high-usage subset of 6.4M employer firms; unverified"),
 "I": (1250, "ASSUMPTION: $5k claimable x 25% fee", 200000, "WEAK: web says ~43% of small businesses sell national brands; dealer count unverified"),
 "J": (2500, "ASSUMPTION: $10k recovered x 25% fee", 111207, "DATA: 111,207 HVAC/plumbing employer establishments (Census CBP 2023)"),
 "K": (750, "ASSUMPTION: $3k recoverable x 25% fee", 430000, "DATA: ~720k restaurants; ASSUMPTION ~60% on delivery apps"),
 "L": (1500, "ASSUMPTION: $100k parcel spend x 5% x 30% fee", 300000, "UNVERIFIED: no shipper count found"),
}

GATES = {  # hard-requirement flags
 "A": "License: none for provider-side appeals, BUT avoid collecting patient balances (collections licensing). HIPAA BAA + security required. Obsolescence: medium risk (payer AI vs provider AI arms race helps demand). Margin: ~60-65% if human review is kept small.",
 "B": "License: FAIL if we collect debts for owners; OK only as owner-name reminders/notices. Obsolescence: HIGH (accounting/invoicing software already ships reminders). Margin OK.",
 "C": "License: none. Obsolescence: medium (Stripe/Shopify shipping native dispute tooling). Margin OK. Purchasing power weak (small merchants).",
 "D": "FAILS obsolescence: commenters already use AI that logs calls into Clio. Reject.",
 "E": "FAILS the 'don't anchor' brief and obsolescence: AI receptionist/follow-up is the most crowded category.",
 "F": "Low willingness to pay; thin evidence.",
 "M": "License: none for filing claims in the carrier's name; VERIFY state commercial-collections rules before any escalation. Not brokering. Obsolescence: medium (TMS/ELD vendors and dedicated detention tools already exist). Evidence caveat: several supporting posts look like idea-validation by builders.",
 "N": "FAILS obsolescence: contractor software already ships change orders and will add AI capture. Behavior problem (owner freezes), not a tooling gap.",
 "G": "No license if audit-only; taking processor residuals would be ISO/sales-agent territory (avoid). Many free/contingent audit competitors. Reddit thread shows owners think fees are small next to other costs.",
 "H": "Almost no Reddit evidence; many established contingency auditors.",
 "I": "Effectively zero Reddit evidence (one removed MSP post). Manufacturer-by-manufacturer rules kill automation.",
 "J": "Only ~10 relevant posts; one venting HVAC tech. Manufacturer portals differ by brand.",
 "K": "Two relevant complaints; platforms auto-refund customers and give ~10 days to dispute.",
 "L": "Real money, but crowded contingency market (25-50% fees) and zero Reddit voice.",
}


def main():
    rows, bad = build_rows()
    if bad:
        print("UNRESOLVED:", *bad, sep="\n  ")
    with open("pain_points.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["cluster", "subreddit", "url", "quote_under_25_words", "industry", "dollar_or_time_impact", "score", "num_comments", "source"])
        w.writerows(rows)
    with open("scores.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["cluster", "status"] + CRIT + ["TOTAL_of_100", "avg_rev_per_client", "clients_for_1M", "reachable_market", "share_needed_pct", "hard_requirement_notes"])
        for name, (v, st) in sorted(CL.items(), key=lambda x: -total(x[1][0])):
            m = MATH[name[0]]
            clients = round(1_000_000 / m[0])
            w.writerow([name, st] + v + [total(v), m[0], clients, m[2], round(100 * clients / m[2], 2), GATES[name[0]]])

    wb = Workbook()
    hdr = PatternFill("solid", fgColor="1F3864"); white = Font(color="FFFFFF", bold=True)
    ws = wb.active; ws.title = "Scores"
    ws.append(["Cluster", "Status"] + CRIT + ["TOTAL /100"])
    for name, (v, st) in sorted(CL.items(), key=lambda x: -total(x[1][0])):
        r = ws.max_row + 1
        ws.append([name, st] + v)
        ws.cell(r, 2 + len(CRIT) + 1, f"=SUM(C{r}:{get_column_letter(1 + len(CRIT) + 1)}{r})+H{r}")
    ws.append([]); ws.append(["Criterion 6 (outreach) is counted twice. Scores are my judgment (1-5), not measured. Edit any cell; totals recalc."])
    ws2 = wb.create_sheet("$1M Math")
    ws2.append(["Cluster", "Avg revenue / client / yr", "Revenue source", "Clients needed for $1M", "Reachable market", "Market source", "Share of market needed", "Flag (>2-3%)", "Hard-requirement notes"])
    for name, (v, st) in sorted(CL.items(), key=lambda x: -total(x[1][0])):
        m = MATH[name[0]]; r = ws2.max_row + 1
        ws2.append([name, m[0], m[1], f"=ROUNDUP(1000000/B{r},0)", m[2], m[3], f"=D{r}/E{r}", f'=IF(G{r}>0.03,"RISK","ok")', GATES[name[0]]])
        ws2.cell(r, 7).number_format = "0.00%"
    ws2.append([]); ws2.append(["Blue-sky inputs are ASSUMPTIONS (change column B / E). Market sizes marked DATA came from web searches, not Reddit."])
    ws3 = wb.create_sheet("Pain Points")
    ws3.append(["Cluster", "Subreddit", "URL", "Quote (<25 words)", "Industry", "Dollar / time impact", "Score", "Comments", "Source"])
    for r in rows:
        ws3.append(r)
    for s in wb.worksheets:
        for c in s[1]:
            c.fill = hdr; c.font = white; c.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "B2"
    for col, wd in zip("ABCDEFGHI", [48, 14, 44, 70, 34, 26, 8, 10, 10]):
        ws3.column_dimensions[col].width = wd
    ws.column_dimensions["A"].width = 60; ws2.column_dimensions["A"].width = 60
    ws2.column_dimensions["C"].width = 50; ws2.column_dimensions["F"].width = 50; ws2.column_dimensions["I"].width = 80
    wb.save("painkiller_workbook.xlsx")
    print(len(rows), "evidence rows;", len(bad), "unresolved")
    for name, (v, st) in sorted(CL.items(), key=lambda x: -total(x[1][0])):
        m = MATH[name[0]]; c = round(1e6 / m[0])
        print(f"{total(v):3d}  {name[:70]:70s} clients={c:5d} share={100*c/m[2]:.2f}%")


if __name__ == "__main__":
    main()
